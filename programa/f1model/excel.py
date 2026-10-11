"""Žaidėjų spėjimai iš totalizatoriaus Excel lentelės. Žaidėjai ir GP pavadinimai – iš žinynų (DB)."""
import openpyxl
import pandas as pd

TYPES = {"Qualifying": "Q", "Race": "R", "Sprint Qualifying": "SQ", "Sprint": "S"}
COLUMNS = ["player", "round", "session", "pos", "predicted"]   # pos – spėta vieta (1–3), predicted – vairuotojas


class ExcelPicks:
    def __init__(self, path, ref):
        self.path, self.players, self.gp_names = path, ref.players, ref.gp_names

    def read(self, events):
        """DataFrame: player, round, session, pos, predicted."""
        if not self.path.exists():
            return pd.DataFrame(columns=COLUMNS)
        wb = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        name_to_round = {}
        for xl, part in self.gp_names.items():
            m = events[events.name.str.contains(part, regex=False)]
            if not m.empty:
                name_to_round[xl] = int(m["round"].iloc[0])
        rows = []
        for player in (p for p in self.players if p in wb.sheetnames):
            gp = typ = None
            for e, f, g, _, i in wb[player].iter_rows(min_row=6, max_row=200, min_col=5, max_col=9, values_only=True):
                gp, typ = e or gp, f or typ
                if g and gp in name_to_round and typ in TYPES:
                    rows.append(dict(player=player, round=name_to_round[gp], session=TYPES[typ],
                                     pos=int(str(g).lstrip("P")), predicted=i))
        return pd.DataFrame(rows, columns=COLUMNS)
